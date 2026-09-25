from pptx import Presentation

prs = Presentation('Proposal_TA_Fahry_Rizky_Samsudin.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides, 1):
    texts = []
    num_pics = 0
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            texts.append(s.text_frame.text.strip().replace('\n', ' -- ')[:60])
        elif s.shape_type == 13: # picture
            num_pics += 1
    head = texts[0] if texts else "NO_TEXT"
    sub = texts[1] if len(texts) > 1 else ""
    print(f"Slide {idx:02d} | Pics: {num_pics} | {head} | {sub}")
