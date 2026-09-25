#!/usr/bin/env python3
"""Create a styled Excel workbook from a compact JSON data specification."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
SCALAR_TYPES = (str, int, float, bool, type(None))


def validate_scalar(value: Any, location: str) -> None:
    if not isinstance(value, SCALAR_TYPES):
        raise ValueError(f"{location} must be a scalar or null, got {type(value).__name__}")


def normalize_sheet(sheet: dict[str, Any]) -> tuple[str, list[str], list[list[Any]]]:
    name = sheet.get("name")
    if not isinstance(name, str) or not name or len(name) > 31 or INVALID_SHEET_CHARS.search(name):
        raise ValueError(f"Invalid Excel sheet name: {name!r}")
    has_columns = "columns" in sheet
    has_rows = "rows" in sheet
    if has_columns == has_rows:
        raise ValueError(f"Sheet {name!r} must contain exactly one of 'columns' or 'rows'")

    if has_columns:
        columns = sheet["columns"]
        if not isinstance(columns, dict) or not columns:
            raise ValueError(f"Sheet {name!r}: 'columns' must be a non-empty object")
        headers = list(columns)
        values = list(columns.values())
        if not all(isinstance(v, list) for v in values):
            raise ValueError(f"Sheet {name!r}: every column value must be an array")
        lengths = {len(v) for v in values}
        if len(lengths) != 1:
            raise ValueError(f"Sheet {name!r}: all columns must have equal lengths")
        rows = [list(row) for row in zip(*values)]
    else:
        source_rows = sheet["rows"]
        if not isinstance(source_rows, list) or not source_rows or not all(isinstance(r, dict) for r in source_rows):
            raise ValueError(f"Sheet {name!r}: 'rows' must be a non-empty array of objects")
        headers = []
        for row in source_rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        rows = [[row.get(header) for header in headers] for row in source_rows]

    if not headers or not all(isinstance(h, str) and h for h in headers):
        raise ValueError(f"Sheet {name!r}: headers must be non-empty strings")
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            validate_scalar(value, f"{name}!{get_column_letter(col_index)}{row_index}")
    return name, headers, rows


def style_table(ws) -> None:
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F1F1F")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_index, column in enumerate(ws.iter_cols(), start=1):
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[get_column_letter(col_index)].width = min(max(max_length + 2, 10), 48)
        for cell in column[1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"


def build_workbook(spec: dict[str, Any]) -> Workbook:
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("'sheets' must be a non-empty array")
    normalized = [normalize_sheet(sheet) for sheet in sheets]
    names = [item[0].casefold() for item in normalized]
    if len(names) != len(set(names)):
        raise ValueError("Sheet names must be unique (case-insensitive)")
    if "readme" in names:
        raise ValueError("'README' is reserved for workbook metadata")

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["field", "value"])
    metadata = spec.get("metadata", {})
    if not isinstance(metadata, dict):
        raise ValueError("'metadata' must be an object")
    defaults = {
        "data_disclaimer": "Example values may be digitized approximations or synthetic look-alikes; consult notes.",
    }
    for key, value in {**defaults, **metadata}.items():
        validate_scalar(value, f"metadata.{key}")
        readme.append([str(key), value])
    style_table(readme)

    for name, headers, rows in normalized:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_table(ws)
    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", required=True, type=Path, help="Input JSON specification")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx path")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".xlsx":
        parser.error("--output must end in .xlsx")
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    workbook = build_workbook(spec)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
